# -*- coding: utf-8 -*-
"""
Exportación Excel/PDF de proyección de caja (solo lectura de datos ya calculados).
No recalcula snapshots ni escribe en BD.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

# ---------------------------------------------------------------------------
# Constantes de presentación (alineadas a proyeccion_caja.CONCEPTOS_ORDEN)
# ---------------------------------------------------------------------------

CONCEPTOS_ORDEN_EXPORT: List[str] = [
    "📥 CxC — Pago Clientes",
    "📤 Proveedores Nacionales",
    "📤 Proveedores Extranjeros",
    "📤 Remuneraciones",
    "📤 Imposiciones AFP/Salud",
    "📤 Impuesto único nómina",
    "📤 Créditos bancarios",
    "📤 Honorarios líquidos",
    "📤 Retenciones 2da categoría",
    "📤 IVA Neto",
    "📤 PPM",
    "📤 F29 total a pagar — código 91",
    "📤 Gastos Aduana/Flete",
    "📤 IVA Importación",
    "📤 Categoría personalizada (cliente)",
]

FILAS_SINTETICAS = frozenset(
    {
        "🏦 Saldo Cartola",
        "💰 Posición Neta Acum.",
        "🧮 Total por día",
    }
)

CONCEPTO_ALIAS_UI = {
    "📥 CxC — Pago Clientes": "📥 Facturas por Cobrar — Clientes",
    "📤 Proveedores Nacionales": "📤 Facturas por Pagar — Proveedores Nacionales",
    "📤 Proveedores Extranjeros": "📤 Facturas por Pagar — Proveedores Extranjeros",
}


def _dec(x: Any) -> Decimal:
    if x is None:
        return Decimal(0)
    if isinstance(x, Decimal):
        return x
    try:
        return Decimal(str(x))
    except Exception:
        return Decimal(0)


def formato_clp(valor: Any) -> str:
    """Formato CLP chileno: $ 5.553.250 (negativos con signo)."""
    n = int(round(float(_dec(valor))))
    sign = "-" if n < 0 else ""
    body = f"{abs(n):,}".replace(",", ".")
    return f"{sign}$ {body}"


def nombre_archivo_seguro(empresa: str, inicio: date, fin: date, ext: str) -> str:
    raw = (empresa or "Empresa").strip() or "Empresa"
    safe = re.sub(r"[^\w\-áéíóúÁÉÍÓÚñÑ]+", "_", raw, flags=re.UNICODE)
    safe = re.sub(r"_+", "_", safe).strip("_")[:60] or "Empresa"
    return f"Flujo_Caja_{safe}_{inicio.isoformat()}_{fin.isoformat()}.{ext.lstrip('.')}"


def concepto_desde_codigo(codigo: str, nombre: str = "") -> str:
    cod = (codigo or "").strip().upper()
    mapping = {
        "CLIENTES": "📥 CxC — Pago Clientes",
        "PROV_NACIONAL": "📤 Proveedores Nacionales",
        "PROV_EXTRANJERO": "📤 Proveedores Extranjeros",
        "REMUNERACIONES": "📤 Remuneraciones",
        "IMPOSICIONES": "📤 Imposiciones AFP/Salud",
        "IU_NOMINA": "📤 Impuesto único nómina",
        "CREDITO_BANCARIO": "📤 Créditos bancarios",
        "HONORARIOS": "📤 Honorarios líquidos",
        "RETENCION": "📤 Retenciones 2da categoría",
        "IVA": "📤 IVA Neto",
        "PPM": "📤 PPM",
        "F29_91": "📤 F29 total a pagar — código 91",
        "GASTOS_IMPORTACION": "📤 Gastos Aduana/Flete",
        "IVA_IMPORTACION": "📤 IVA Importación",
    }
    if cod in mapping:
        return mapping[cod]
    if cod:
        return "📤 Categoría personalizada (cliente)"
    if "cliente" in (nombre or "").lower():
        return "📥 CxC — Pago Clientes"
    return "📤 Categoría personalizada (cliente)"


# ---------------------------------------------------------------------------
# DTO
# ---------------------------------------------------------------------------


@dataclass
class LineaExportacion:
    fecha: date
    concepto: str
    categoria: str
    descripcion: str
    monto: Decimal
    confianza: str
    origen: str


@dataclass
class FacturaExportacion:
    tipo: str  # por_cobrar / por_pagar
    fecha_vencimiento: Optional[date]
    razon_social: str
    folio: str
    monto: Decimal
    saldo: Decimal
    estado: str
    tipo_confianza: str


@dataclass
class ParametrosExportacion:
    tasa_ppm: Optional[Decimal] = None
    tasa_retencion_honorarios: Optional[Decimal] = None
    dia_pago_impuestos: Optional[int] = None
    dia_pago_remuneraciones: Optional[int] = None
    dia_pago_imposiciones: Optional[int] = None
    venta_global_esperada_mes: Optional[Decimal] = None
    porcentaje_ventas_contado: Optional[Decimal] = None
    compra_global_esperada_mes: Optional[Decimal] = None
    porcentaje_compras_contado: Optional[Decimal] = None
    porcentaje_morosidad_cxc: Optional[Decimal] = None
    porcentaje_recuperabilidad_morosos: Optional[Decimal] = None
    periodo_fuente_remuneraciones: Optional[str] = None


@dataclass
class ContextoExportacionProyeccion:
    """Contexto aislado por usuario y snapshot (solo lectura)."""

    user_id: int
    snapshot_user_id: int
    empresa: str
    fecha_emision: date
    snapshot_id: int
    version: int
    etiqueta: str
    notas: str
    periodo_inicio: date
    periodo_fin: date
    saldo_inicial: Decimal
    cobros_cxc: Decimal
    egresos: Decimal
    flujo_neto: Decimal
    saldo_final: Decimal
    pct_estimado_manual: float
    advertencias: List[str] = field(default_factory=list)
    lineas: List[LineaExportacion] = field(default_factory=list)
    facturas: List[FacturaExportacion] = field(default_factory=list)
    parametros: ParametrosExportacion = field(default_factory=ParametrosExportacion)

    def validar_aislamiento(self) -> None:
        if int(self.user_id) != int(self.snapshot_user_id):
            raise PermissionError(
                "Exportación denegada: el snapshot no pertenece al usuario de la sesión."
            )


# ---------------------------------------------------------------------------
# Construcción de tablas derivadas
# ---------------------------------------------------------------------------


def construir_flujo_diario(
    lineas: Sequence[LineaExportacion],
    saldo_inicial: Decimal,
) -> pd.DataFrame:
    if not lineas:
        return pd.DataFrame(
            columns=["fecha", "ingresos", "egresos", "neto_diario", "saldo_acumulado"]
        )
    rows: Dict[date, Dict[str, Decimal]] = {}
    for ln in lineas:
        bucket = rows.setdefault(
            ln.fecha, {"ingresos": Decimal(0), "egresos": Decimal(0)}
        )
        if ln.monto >= 0:
            bucket["ingresos"] += ln.monto
        else:
            bucket["egresos"] += ln.monto
    out = []
    acum = _dec(saldo_inicial)
    for fd in sorted(rows.keys()):
        ing = rows[fd]["ingresos"]
        egr = rows[fd]["egresos"]
        neto = ing + egr
        acum += neto
        out.append(
            {
                "fecha": fd,
                "ingresos": float(ing),
                "egresos": float(egr),
                "neto_diario": float(neto),
                "saldo_acumulado": float(acum),
            }
        )
    return pd.DataFrame(out)


def construir_matriz_concepto_fecha(lineas: Sequence[LineaExportacion]) -> pd.DataFrame:
    if not lineas:
        return pd.DataFrame(columns=["Concepto", "TOTAL"])
    df = pd.DataFrame(
        [
            {
                "concepto": ln.concepto,
                "fecha": ln.fecha,
                "monto": float(ln.monto),
            }
            for ln in lineas
            if ln.concepto not in FILAS_SINTETICAS
        ]
    )
    if df.empty:
        return pd.DataFrame(columns=["Concepto", "TOTAL"])
    df["fecha_s"] = pd.to_datetime(df["fecha"]).dt.strftime("%Y-%m-%d")
    pivot = (
        df.groupby(["concepto", "fecha_s"], as_index=False)["monto"]
        .sum()
        .pivot(index="concepto", columns="fecha_s", values="monto")
        .fillna(0.0)
    )
    # Orden preferente + conceptos extra
    orden = [c for c in CONCEPTOS_ORDEN_EXPORT if c in pivot.index]
    extra = [c for c in pivot.index if c not in orden]
    pivot = pivot.loc[orden + sorted(extra)]
    # Omitir conceptos completamente en cero
    pivot = pivot.loc[(pivot.abs().sum(axis=1) > 0)]
    pivot["TOTAL"] = pivot.sum(axis=1)
    pivot = pivot.reset_index().rename(columns={"concepto": "Concepto"})
    return pivot


def construir_df_detalle(lineas: Sequence[LineaExportacion]) -> pd.DataFrame:
    rows = [
        {
            "fecha": ln.fecha,
            "concepto": ln.concepto,
            "categoría": ln.categoria,
            "descripción": ln.descripcion,
            "monto": float(ln.monto),
            "confianza": ln.confianza,
            "origen": ln.origen,
        }
        for ln in sorted(lineas, key=lambda x: (x.fecha, x.concepto, x.descripcion))
    ]
    return pd.DataFrame(
        rows,
        columns=["fecha", "concepto", "categoría", "descripción", "monto", "confianza", "origen"],
    )


def construir_df_facturas(
    facturas: Sequence[FacturaExportacion], tipo: str
) -> pd.DataFrame:
    rows = [
        {
            "fecha_vencimiento": f.fecha_vencimiento,
            "razon_social": f.razon_social,
            "folio": f.folio,
            "monto": float(f.monto),
            "saldo": float(f.saldo),
            "estado": f.estado,
            "confianza": f.tipo_confianza,
        }
        for f in facturas
        if (f.tipo or "").strip().lower() == tipo
    ]
    return pd.DataFrame(
        rows,
        columns=[
            "fecha_vencimiento",
            "razon_social",
            "folio",
            "monto",
            "saldo",
            "estado",
            "confianza",
        ],
    )


def totales_por_concepto(lineas: Sequence[LineaExportacion]) -> List[Tuple[str, float]]:
    agg: Dict[str, Decimal] = {}
    for ln in lineas:
        if ln.concepto in FILAS_SINTETICAS:
            continue
        agg[ln.concepto] = agg.get(ln.concepto, Decimal(0)) + ln.monto
    items = [(k, float(v)) for k, v in agg.items() if abs(v) > 0]
    order_idx = {c: i for i, c in enumerate(CONCEPTOS_ORDEN_EXPORT)}
    items.sort(key=lambda kv: (order_idx.get(kv[0], 999), kv[0]))
    return items


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

_FILL_TITLE = PatternFill("solid", fgColor="0E5A8A")
_FILL_HEADER = PatternFill("solid", fgColor="1D8F6E")
_FILL_TOTAL = PatternFill("solid", fgColor="E8F5F0")
_FILL_KPI = PatternFill("solid", fgColor="F7FAF9")
_FONT_TITLE = Font(name="Calibri", bold=True, color="FFFFFF", size=14)
_FONT_HEADER = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_FONT_BOLD = Font(name="Calibri", bold=True, size=11)
_FONT_NEG = Font(name="Calibri", color="C00000")
_THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
_NUM_CLP = '#,##0;[Red]-#,##0'


def _style_header_row(ws, row: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _FILL_HEADER
        cell.font = _FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN


def _autosize(ws, min_w: int = 10, max_w: int = 42) -> None:
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        length = 0
        for cell in col_cells:
            if cell.value is None:
                continue
            length = max(length, min(len(str(cell.value)), max_w))
        ws.column_dimensions[letter].width = max(min_w, length + 2)


def _print_setup(ws, landscape: bool = True) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.print_title_rows = "1:2"


def _write_dataframe_sheet(
    ws,
    df: pd.DataFrame,
    *,
    title: str,
    money_cols: Sequence[str],
    date_cols: Sequence[str],
    freeze: str = "A3",
) -> None:
    ws["A1"] = title
    ws["A1"].font = _FONT_TITLE
    ws["A1"].fill = _FILL_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(df.columns), 1))

    if df.empty:
        ws["A3"] = "Sin datos en el período seleccionado."
        _print_setup(ws)
        return

    header_row = 2
    for c_idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=c_idx, value=str(col))
        cell.fill = _FILL_HEADER
        cell.font = _FONT_HEADER
        cell.border = _THIN
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    money_set = set(money_cols)
    date_set = set(date_cols)
    for r_idx, row in enumerate(df.itertuples(index=False), start=3):
        for c_idx, (col_name, value) in enumerate(zip(df.columns, row), start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = _THIN
            if col_name in date_set and value is not None and not (isinstance(value, float) and pd.isna(value)):
                if isinstance(value, datetime):
                    cell.value = value.date()
                elif isinstance(value, date):
                    cell.value = value
                else:
                    try:
                        cell.value = pd.to_datetime(value).date()
                    except Exception:
                        cell.value = value
                cell.number_format = "DD/MM/YYYY"
            elif col_name in money_set:
                try:
                    cell.value = float(value) if value is not None and not pd.isna(value) else 0.0
                except Exception:
                    cell.value = 0.0
                cell.number_format = _NUM_CLP
                if isinstance(cell.value, (int, float)) and cell.value < 0:
                    cell.font = _FONT_NEG
            else:
                cell.value = None if (isinstance(value, float) and pd.isna(value)) else value

    # Resaltar fila TOTAL si existe
    for r_idx in range(3, ws.max_row + 1):
        first = ws.cell(row=r_idx, column=1).value
        if first is not None and str(first).strip().upper() in {"TOTAL", "TOTALES"}:
            for c_idx in range(1, ws.max_column + 1):
                ws.cell(row=r_idx, column=c_idx).fill = _FILL_TOTAL
                ws.cell(row=r_idx, column=c_idx).font = _FONT_BOLD

    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.freeze_panes = freeze
    _autosize(ws)
    _print_setup(ws, landscape=len(df.columns) > 5)


def generar_excel_bytes(ctx: ContextoExportacionProyeccion) -> bytes:
    ctx.validar_aislamiento()
    wb = Workbook()

    # --- Resumen ---
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Informe de Flujo de Caja Proyectado"
    ws["A1"].font = _FONT_TITLE
    ws["A1"].fill = _FILL_TITLE
    ws.merge_cells("A1:B1")

    resumen_rows = [
        ("Empresa", ctx.empresa),
        ("Fecha de emisión", ctx.fecha_emision),
        ("Snapshot ID", ctx.snapshot_id),
        ("Versión", ctx.version),
        ("Etiqueta", ctx.etiqueta or "—"),
        ("Horizonte inicio", ctx.periodo_inicio),
        ("Horizonte fin", ctx.periodo_fin),
        ("Saldo inicial", float(ctx.saldo_inicial)),
        ("Cobros CxC", float(ctx.cobros_cxc)),
        ("Egresos", float(ctx.egresos)),
        ("Flujo neto", float(ctx.flujo_neto)),
        ("Saldo final", float(ctx.saldo_final)),
        ("% estimado/manual", float(ctx.pct_estimado_manual)),
        ("Notas", ctx.notas or "—"),
    ]
    ws["A2"] = "Campo"
    ws["B2"] = "Valor"
    _style_header_row(ws, 2, 2)
    money_labels = {
        "Saldo inicial",
        "Cobros CxC",
        "Egresos",
        "Flujo neto",
        "Saldo final",
    }
    date_labels = {"Fecha de emisión", "Horizonte inicio", "Horizonte fin"}
    for i, (k, v) in enumerate(resumen_rows, start=3):
        ws.cell(row=i, column=1, value=k).font = _FONT_BOLD
        ws.cell(row=i, column=1).fill = _FILL_KPI
        ws.cell(row=i, column=1).border = _THIN
        cell = ws.cell(row=i, column=2, value=v)
        cell.border = _THIN
        if k in money_labels:
            cell.value = float(v)
            cell.number_format = _NUM_CLP
            if float(v) < 0:
                cell.font = _FONT_NEG
        elif k in date_labels and isinstance(v, date):
            cell.number_format = "DD/MM/YYYY"
        elif k == "% estimado/manual":
            cell.value = float(v)
            cell.number_format = "0.0%"

    start_adv = 3 + len(resumen_rows) + 1
    ws.cell(row=start_adv, column=1, value="Advertencias").font = _FONT_BOLD
    if ctx.advertencias:
        for j, adv in enumerate(ctx.advertencias):
            ws.cell(row=start_adv + 1 + j, column=1, value=str(adv))
            ws.merge_cells(
                start_row=start_adv + 1 + j,
                start_column=1,
                end_row=start_adv + 1 + j,
                end_column=2,
            )
    else:
        ws.cell(row=start_adv + 1, column=1, value="Sin advertencias relevantes.")
    ws.freeze_panes = "A3"
    _autosize(ws, max_w=80)
    _print_setup(ws, landscape=False)

    # --- Flujo diario ---
    df_flujo = construir_flujo_diario(ctx.lineas, ctx.saldo_inicial)
    ws_f = wb.create_sheet("Flujo diario")
    ws_f["A1"] = "Flujo diario proyectado"
    ws_f["A1"].font = _FONT_TITLE
    ws_f["A1"].fill = _FILL_TITLE
    ws_f.merge_cells("A1:E1")
    ws_f["A2"] = "Saldo inicial de caja"
    ws_f["A2"].font = _FONT_BOLD
    ws_f["A2"].fill = _FILL_KPI
    ws_f["A2"].border = _THIN
    ws_f["B2"] = float(ctx.saldo_inicial)
    ws_f["B2"].number_format = _NUM_CLP
    ws_f["B2"].font = _FONT_BOLD
    ws_f["B2"].fill = _FILL_KPI
    ws_f["B2"].border = _THIN
    if float(ctx.saldo_inicial) < 0:
        ws_f["B2"].font = Font(name="Calibri", bold=True, color="C00000", size=11)
    # Tabla desde fila 4 (fila 3 = encabezados); saldo_acumulado sin cambiar de fórmula.
    header_row = 3
    data_start = 4
    if df_flujo.empty:
        ws_f["A4"] = "Sin datos en el período seleccionado."
    else:
        cols = list(df_flujo.columns)
        for c_idx, col in enumerate(cols, start=1):
            cell = ws_f.cell(row=header_row, column=c_idx, value=str(col))
            cell.fill = _FILL_HEADER
            cell.font = _FONT_HEADER
            cell.border = _THIN
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        money_set = {"ingresos", "egresos", "neto_diario", "saldo_acumulado"}
        for r_idx, row in enumerate(df_flujo.itertuples(index=False), start=data_start):
            for c_idx, (col_name, value) in enumerate(zip(cols, row), start=1):
                cell = ws_f.cell(row=r_idx, column=c_idx)
                cell.border = _THIN
                if col_name == "fecha":
                    if isinstance(value, datetime):
                        cell.value = value.date()
                    elif isinstance(value, date):
                        cell.value = value
                    else:
                        try:
                            cell.value = pd.to_datetime(value).date()
                        except Exception:
                            cell.value = value
                    cell.number_format = "DD/MM/YYYY"
                elif col_name in money_set:
                    cell.value = float(value) if value is not None and not pd.isna(value) else 0.0
                    cell.number_format = _NUM_CLP
                    if isinstance(cell.value, (int, float)) and cell.value < 0:
                        cell.font = _FONT_NEG
                else:
                    cell.value = value
        ws_f.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{ws_f.max_row}"
    ws_f.freeze_panes = "A4"
    _autosize(ws_f)
    _print_setup(ws_f, landscape=True)

    # --- Detalle por concepto (matriz) ---
    df_mat = construir_matriz_concepto_fecha(ctx.lineas)
    ws_m = wb.create_sheet("Detalle por concepto")
    money_m = [c for c in df_mat.columns if c != "Concepto"]
    _write_dataframe_sheet(
        ws_m,
        df_mat,
        title="Matriz concepto × fecha (sin filas sintéticas ni conceptos en cero)",
        money_cols=money_m,
        date_cols=[],
        freeze="B3",
    )
    # Destacar columna TOTAL
    if not df_mat.empty and "TOTAL" in df_mat.columns:
        tot_col = list(df_mat.columns).index("TOTAL") + 1
        for r in range(3, ws_m.max_row + 1):
            ws_m.cell(row=r, column=tot_col).fill = _FILL_TOTAL
            ws_m.cell(row=r, column=tot_col).font = _FONT_BOLD

    # --- Detalle ---
    df_det = construir_df_detalle(ctx.lineas)
    ws_d = wb.create_sheet("Detalle")
    _write_dataframe_sheet(
        ws_d,
        df_det,
        title="Detalle de líneas del período visible",
        money_cols=["monto"],
        date_cols=["fecha"],
    )

    # --- CxC / CxP ---
    df_cxc = construir_df_facturas(ctx.facturas, "por_cobrar")
    df_cxp = construir_df_facturas(ctx.facturas, "por_pagar")
    ws_cxc = wb.create_sheet("CxC")
    _write_dataframe_sheet(
        ws_cxc,
        df_cxc,
        title="Facturas por cobrar (última carga del usuario)",
        money_cols=["monto", "saldo"],
        date_cols=["fecha_vencimiento"],
    )
    ws_cxp = wb.create_sheet("CxP")
    _write_dataframe_sheet(
        ws_cxp,
        df_cxp,
        title="Facturas por pagar (última carga del usuario)",
        money_cols=["monto", "saldo"],
        date_cols=["fecha_vencimiento"],
    )

    # --- Supuestos ---
    ws_s = wb.create_sheet("Supuestos")
    ws_s["A1"] = "Parámetros y supuestos del escenario"
    ws_s["A1"].font = _FONT_TITLE
    ws_s["A1"].fill = _FILL_TITLE
    ws_s.merge_cells("A1:B1")
    ws_s["A2"] = "Parámetro"
    ws_s["B2"] = "Valor"
    _style_header_row(ws_s, 2, 2)
    p = ctx.parametros

    def _pct(v: Optional[Decimal]) -> Any:
        if v is None:
            return "—"
        return float(v)

    supuestos = [
        ("Tasa PPM", _pct(p.tasa_ppm), "pct"),
        ("Tasa retención honorarios", _pct(p.tasa_retencion_honorarios), "pct"),
        ("Día pago impuestos (F29)", p.dia_pago_impuestos if p.dia_pago_impuestos is not None else "—", "txt"),
        ("Día pago remuneraciones", p.dia_pago_remuneraciones if p.dia_pago_remuneraciones is not None else "—", "txt"),
        ("Día pago imposiciones", p.dia_pago_imposiciones if p.dia_pago_imposiciones is not None else "—", "txt"),
        ("Venta global esperada / mes", float(p.venta_global_esperada_mes or 0) if p.venta_global_esperada_mes is not None else "—", "money"),
        ("% ventas contado", _pct(p.porcentaje_ventas_contado), "pct"),
        ("Compra global esperada / mes", float(p.compra_global_esperada_mes or 0) if p.compra_global_esperada_mes is not None else "—", "money"),
        ("% compras contado", _pct(p.porcentaje_compras_contado), "pct"),
        ("% morosidad CxC", _pct(p.porcentaje_morosidad_cxc), "pct"),
        ("% recuperabilidad morosos", _pct(p.porcentaje_recuperabilidad_morosos), "pct"),
        (
            "Período fuente libro remuneraciones",
            p.periodo_fuente_remuneraciones or "No detectado en el snapshot visible",
            "txt",
        ),
    ]
    for i, (lab, val, kind) in enumerate(supuestos, start=3):
        ws_s.cell(row=i, column=1, value=lab).font = _FONT_BOLD
        ws_s.cell(row=i, column=1).border = _THIN
        cell = ws_s.cell(row=i, column=2, value=val)
        cell.border = _THIN
        if kind == "money" and isinstance(val, (int, float)):
            cell.number_format = _NUM_CLP
        elif kind == "pct" and isinstance(val, (int, float)):
            cell.number_format = "0.00%"
    # Notas de supuestos comerciales detectados en líneas
    row_n = 3 + len(supuestos) + 1
    ws_s.cell(row=row_n, column=1, value="Notas del escenario").font = _FONT_BOLD
    ws_s.cell(row=row_n + 1, column=1, value=ctx.notas or "—")
    ws_s.merge_cells(start_row=row_n + 1, start_column=1, end_row=row_n + 1, end_column=2)
    row_a = row_n + 3
    ws_s.cell(row=row_a, column=1, value="Advertencias").font = _FONT_BOLD
    if ctx.advertencias:
        for j, adv in enumerate(ctx.advertencias):
            ws_s.cell(row=row_a + 1 + j, column=1, value=str(adv))
            ws_s.merge_cells(
                start_row=row_a + 1 + j,
                start_column=1,
                end_row=row_a + 1 + j,
                end_column=2,
            )
    else:
        ws_s.cell(row=row_a + 1, column=1, value="Sin advertencias.")
    ws_s.freeze_panes = "A3"
    _autosize(ws_s, max_w=70)
    _print_setup(ws_s, landscape=False)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# PDF (ReportLab)
# ---------------------------------------------------------------------------


def generar_pdf_bytes(ctx: ContextoExportacionProyeccion) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.graphics.charts.barcharts import VerticalBarChart

    ctx.validar_aislamiento()
    buf = io.BytesIO()
    page_w, page_h = A4

    styles = getSampleStyleSheet()
    style_title = ParagraphStyle(
        "TituloExp",
        parent=styles["Heading1"],
        fontSize=16,
        textColor=colors.HexColor("#0E5A8A"),
        spaceAfter=6,
        alignment=TA_CENTER,
        leading=20,
    )
    style_h2 = ParagraphStyle(
        "H2Exp",
        parent=styles["Heading2"],
        fontSize=12,
        textColor=colors.HexColor("#123729"),
        spaceBefore=8,
        spaceAfter=6,
        leading=15,
    )
    style_body = ParagraphStyle(
        "BodyExp",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#222222"),
    )
    style_small = ParagraphStyle(
        "SmallExp",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#555555"),
    )
    style_kpi = ParagraphStyle(
        "KpiExp",
        parent=styles["Normal"],
        fontSize=9,
        leading=11,
        alignment=TA_CENTER,
        textColor=colors.white,
    )

    empresa = ctx.empresa or "Empresa"

    def _header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#0E5A8A"))
        canvas.drawString(1.5 * cm, page_h - 1.1 * cm, f"{empresa} — Flujo de caja proyectado")
        canvas.setStrokeColor(colors.HexColor("#1D8F6E"))
        canvas.setLineWidth(0.8)
        canvas.line(1.5 * cm, page_h - 1.25 * cm, page_w - 1.5 * cm, page_h - 1.25 * cm)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(
            1.5 * cm,
            1.0 * cm,
            f"Emisión {ctx.fecha_emision.strftime('%d/%m/%Y')} · Snapshot v{ctx.version} (id {ctx.snapshot_id})",
        )
        canvas.drawRightString(page_w - 1.5 * cm, 1.0 * cm, f"Página {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.6 * cm,
        title=f"Flujo de Caja — {empresa}",
        author=empresa,
    )
    story: List[Any] = []

    story.append(Paragraph("Informe de Flujo de Caja Proyectado", style_title))
    story.append(Spacer(1, 4))
    story.append(
        Paragraph(
            f"<b>Empresa:</b> {empresa}<br/>"
            f"<b>Fecha de emisión:</b> {ctx.fecha_emision.strftime('%d/%m/%Y')}<br/>"
            f"<b>Horizonte:</b> {ctx.periodo_inicio.strftime('%d/%m/%Y')} → {ctx.periodo_fin.strftime('%d/%m/%Y')}<br/>"
            f"<b>Snapshot:</b> v{ctx.version} · {(ctx.etiqueta or 'sin etiqueta')} · id {ctx.snapshot_id}",
            style_body,
        )
    )
    story.append(Spacer(1, 10))

    # Tarjetas KPI
    kpi_data = [
        [
            Paragraph(f"<b>Saldo inicial</b><br/>{formato_clp(ctx.saldo_inicial)}", style_kpi),
            Paragraph(f"<b>Cobros CxC</b><br/>{formato_clp(ctx.cobros_cxc)}", style_kpi),
            Paragraph(f"<b>Egresos</b><br/>{formato_clp(ctx.egresos)}", style_kpi),
        ],
        [
            Paragraph(f"<b>Flujo neto</b><br/>{formato_clp(ctx.flujo_neto)}", style_kpi),
            Paragraph(f"<b>Saldo final</b><br/>{formato_clp(ctx.saldo_final)}", style_kpi),
            Paragraph(
                f"<b>% estimado/manual</b><br/>{ctx.pct_estimado_manual * 100:.1f}%",
                style_kpi,
            ),
        ],
    ]
    kpi_table = Table(kpi_data, colWidths=[5.5 * cm, 5.5 * cm, 5.5 * cm], hAlign="CENTER")
    kpi_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#2E4C97")),
                ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#0EA35A")),
                ("BACKGROUND", (2, 0), (2, 0), colors.HexColor("#B6382E")),
                ("BACKGROUND", (0, 1), (0, 1), colors.HexColor("#4D6CD9")),
                ("BACKGROUND", (1, 1), (1, 1), colors.HexColor("#1D8F6E")),
                ("BACKGROUND", (2, 1), (2, 1), colors.HexColor("#9B8326")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.white),
                ("INNERGRID", (0, 0), (-1, -1), 2, colors.white),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(kpi_table)
    story.append(Spacer(1, 8))
    story.append(
        Paragraph(
            f"Verificación: saldo inicial + flujo neto = saldo final → "
            f"{formato_clp(ctx.saldo_inicial)} + {formato_clp(ctx.flujo_neto)} = {formato_clp(ctx.saldo_final)}.",
            style_small,
        )
    )
    if ctx.advertencias:
        story.append(Spacer(1, 6))
        story.append(Paragraph("<b>Advertencias</b>", style_h2))
        for adv in ctx.advertencias[:12]:
            story.append(Paragraph(f"• {adv}", style_small))

    # Página 2 — evolución diaria
    story.append(PageBreak())
    story.append(Paragraph("Evolución diaria del flujo", style_h2))
    df_flujo = construir_flujo_diario(ctx.lineas, ctx.saldo_inicial)

    if df_flujo.empty:
        story.append(Paragraph("Sin movimientos en el período visible.", style_body))
    else:
        # Tabla (máximo razonable; si hay muchas fechas, resumir últimas/primeras)
        max_rows = 28
        df_show = df_flujo
        if len(df_flujo) > max_rows:
            story.append(
                Paragraph(
                    f"Se muestran las primeras {max_rows} fechas de {len(df_flujo)} días "
                    f"(detalle completo en Excel).",
                    style_small,
                )
            )
            df_show = df_flujo.head(max_rows)

        header = ["Fecha", "Ingresos", "Egresos", "Neto", "Saldo acum."]
        table_data = [header]
        for _, r in df_show.iterrows():
            fd = r["fecha"]
            fd_s = fd.strftime("%d/%m/%Y") if hasattr(fd, "strftime") else str(fd)
            table_data.append(
                [
                    fd_s,
                    formato_clp(r["ingresos"]),
                    formato_clp(r["egresos"]),
                    formato_clp(r["neto_diario"]),
                    formato_clp(r["saldo_acumulado"]),
                ]
            )
        t = Table(
            table_data,
            colWidths=[2.8 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm, 3.6 * cm],
            repeatRows=1,
        )
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D8F6E")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                    ("ALIGN", (0, 0), (0, -1), "CENTER"),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAF9")]),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story.append(t)
        story.append(Spacer(1, 10))

        # Gráfico simple ReportLab (neto diario)
        chart_src = df_flujo.head(20)
        drawing = Drawing(460, 180)
        chart = VerticalBarChart()
        chart.x = 40
        chart.y = 30
        chart.height = 130
        chart.width = 400
        chart.data = [list(chart_src["neto_diario"].astype(float))]
        chart.categoryAxis.categoryNames = [
            (d.strftime("%d/%m") if hasattr(d, "strftime") else str(d)[5:10])
            for d in chart_src["fecha"]
        ]
        chart.bars[0].fillColor = colors.HexColor("#0E5A8A")
        chart.valueAxis.labels.fontSize = 7
        chart.categoryAxis.labels.fontSize = 6
        chart.categoryAxis.labels.angle = 45
        chart.categoryAxis.labels.boxAnchor = "ne"
        drawing.add(chart)
        drawing.add(
            String(120, 165, "Neto diario (primeras fechas del horizonte)", fontSize=9)
        )
        story.append(drawing)

    # Página 3 — totales por concepto y supuestos
    story.append(PageBreak())
    story.append(Paragraph("Totales por concepto", style_h2))
    tot = totales_por_concepto(ctx.lineas)
    if not tot:
        story.append(Paragraph("Sin conceptos con monto distinto de cero.", style_body))
    else:
        td = [["Concepto", "Total", "Confianza / origen predominante"]]
        # confianza/origen predominante por concepto
        conf_map: Dict[str, Dict[str, Decimal]] = {}
        orig_map: Dict[str, Dict[str, Decimal]] = {}
        for ln in ctx.lineas:
            conf_map.setdefault(ln.concepto, {})
            orig_map.setdefault(ln.concepto, {})
            conf_map[ln.concepto][ln.confianza or "—"] = conf_map[ln.concepto].get(
                ln.confianza or "—", Decimal(0)
            ) + abs(ln.monto)
            orig_map[ln.concepto][ln.origen or "—"] = orig_map[ln.concepto].get(
                ln.origen or "—", Decimal(0)
            ) + abs(ln.monto)

        def _dom(d: Mapping[str, Decimal]) -> str:
            if not d:
                return "—"
            return max(d.items(), key=lambda kv: kv[1])[0]

        for concepto, monto in tot:
            td.append(
                [
                    Paragraph(concepto, style_small),
                    formato_clp(monto),
                    f"{_dom(conf_map.get(concepto, {}))} / {_dom(orig_map.get(concepto, {}))}",
                ]
            )
        t2 = Table(td, colWidths=[8.5 * cm, 3.5 * cm, 5 * cm], repeatRows=1)
        t2.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0E5A8A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("ALIGN", (1, 1), (1, -1), "RIGHT"),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FBFF")]),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(t2)

    story.append(Spacer(1, 10))
    story.append(Paragraph("Supuestos principales", style_h2))
    p = ctx.parametros

    def _pct_txt(v: Optional[Decimal]) -> str:
        if v is None:
            return "—"
        return f"{float(v) * 100:.2f}%"

    supuestos_txt = [
        f"Días de pago: impuestos {p.dia_pago_impuestos or '—'}, "
        f"remuneraciones {p.dia_pago_remuneraciones or '—'}, "
        f"imposiciones {p.dia_pago_imposiciones or '—'}.",
        f"Tasa PPM: {_pct_txt(p.tasa_ppm)} · Retención honorarios: {_pct_txt(p.tasa_retencion_honorarios)}.",
        f"Morosidad CxC: {_pct_txt(p.porcentaje_morosidad_cxc)} · "
        f"Recuperabilidad: {_pct_txt(p.porcentaje_recuperabilidad_morosos)}.",
        f"Ventas contado: {_pct_txt(p.porcentaje_ventas_contado)} · "
        f"Compras contado: {_pct_txt(p.porcentaje_compras_contado)}.",
        f"Período fuente remuneraciones: {p.periodo_fuente_remuneraciones or 'No detectado'}.",
        f"% estimado/manual del flujo: {ctx.pct_estimado_manual * 100:.1f}%.",
    ]
    for line in supuestos_txt:
        story.append(Paragraph(f"• {line}", style_body))

    if ctx.notas:
        story.append(Spacer(1, 6))
        story.append(Paragraph("<b>Notas del escenario</b>", style_body))
        story.append(Paragraph(ctx.notas, style_small))

    doc.build(story, onFirstPage=_header_footer, onLaterPages=_header_footer)
    return buf.getvalue()


def exportar_excel_y_pdf(
    ctx: ContextoExportacionProyeccion,
) -> Tuple[bytes, bytes, str, str]:
    """Retorna (excel_bytes, pdf_bytes, nombre_xlsx, nombre_pdf)."""
    ctx.validar_aislamiento()
    xlsx = generar_excel_bytes(ctx)
    pdf = generar_pdf_bytes(ctx)
    n_xlsx = nombre_archivo_seguro(ctx.empresa, ctx.periodo_inicio, ctx.periodo_fin, "xlsx")
    n_pdf = nombre_archivo_seguro(ctx.empresa, ctx.periodo_inicio, ctx.periodo_fin, "pdf")
    return xlsx, pdf, n_xlsx, n_pdf
